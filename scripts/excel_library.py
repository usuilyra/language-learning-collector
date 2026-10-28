"""Deterministic workbook operations for the Language Learning Collector skill.

Public functions return dictionaries with an ``ok`` boolean. Expected input,
validation, lookup, and persistence failures are data, not raised exceptions.
"""

from __future__ import annotations

import json
import os
import re
import shutil
import sys
import unicodedata
import uuid
from pathlib import Path
from typing import Any, Callable, Iterable, Mapping, Sequence

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Border, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.cell_range import CellRange, MultiCellRange
    from openpyxl.worksheet.datavalidation import DataValidation
except ModuleNotFoundError as exc:
    if exc.name != "openpyxl":
        raise

    requirements_path = Path(__file__).resolve().parent.parent / "requirements.txt"
    if not requirements_path.is_file():
        raise ModuleNotFoundError(
            "The required package 'openpyxl' is missing, and requirements.txt "
            "was not found in the Skill folder. Reinstall the complete Skill folder."
        ) from exc

    install_command = f'"{sys.executable}" -m pip install -r "{requirements_path}"'
    raise ModuleNotFoundError(
        "The required package 'openpyxl' is not installed for this Python "
        f"interpreter. Ask the user for permission before running: {install_command}"
    ) from exc


SHEETS = ("Library", "Vocabulary", "Grammar", "Expression")
KNOWLEDGE_SHEETS = ("Vocabulary", "Grammar", "Expression")
HEADERS = {
    "Library": ("language",),
    "Vocabulary": (
        "id",
        "item",
        "meaning",
        "pos",
        "collocations",
        "example",
        "example_translation",
        "frequency",
        "mastery",
    ),
    "Grammar": (
        "id",
        "item",
        "meaning",
        "usage",
        "example",
        "example_translation",
        "frequency",
        "mastery",
    ),
    "Expression": (
        "id",
        "item",
        "meaning",
        "usage",
        "example",
        "example_translation",
        "frequency",
        "mastery",
    ),
}
REQUIRED_TEXT = {
    "Vocabulary": ("id", "item", "meaning"),
    "Grammar": ("id", "item", "meaning", "usage"),
    "Expression": ("id", "item", "meaning", "usage"),
}
OPTIONAL_TEXT = {
    "Vocabulary": ("pos", "collocations", "example", "example_translation"),
    "Grammar": ("example", "example_translation"),
    "Expression": ("example", "example_translation"),
}
UPDATE_FIELDS = {
    "Vocabulary": {"meaning", "pos", "collocations", "example", "example_translation", "mastery"},
    "Grammar": {"meaning", "usage", "example", "example_translation", "mastery"},
    "Expression": {"meaning", "usage", "example", "example_translation", "mastery"},
}
_MASTERY_OPTIONS = ("unmastered", "unfamiliar", "mastered")
MASTERY_VALUES = set(_MASTERY_OPTIONS)
DEFAULT_REGISTRY_PATH = Path.home() / ".language-learning-collector" / "workbooks.json"
_LANGUAGE_KEY = re.compile(r"^[a-z][a-z0-9_-]*$")
_HEADER_FILL = PatternFill(fill_type="solid", fgColor="FFDDEBF7")
_BORDER_SIDE = Side(style="thin", color="FF000000")
_CELL_BORDER = Border(left=_BORDER_SIDE, right=_BORDER_SIDE, top=_BORDER_SIDE, bottom=_BORDER_SIDE)


def normalize_item(value: str) -> str:
    """Apply the schema's deterministic item normalization."""

    if not isinstance(value, str):
        raise TypeError("item must be text")
    return " ".join(unicodedata.normalize("NFC", value).strip().split())


def _ok(**values: Any) -> dict[str, Any]:
    return {"ok": True, **values}


def _fail(reason: str, detail: Any, **values: Any) -> dict[str, Any]:
    return {"ok": False, "reason": reason, "detail": detail, **values}


def _issue(sheet: str | None, rule: str, reason: str, row: int | None = None, **detail: Any) -> dict[str, Any]:
    issue = {"sheet": sheet, "row": row, "rule": rule, "reason": reason}
    issue.update(detail)
    return issue


def _language_issue(language: Any) -> dict[str, Any] | None:
    if not isinstance(language, str) or not _LANGUAGE_KEY.fullmatch(language):
        return _issue(
            "Library",
            "canonical_language_key",
            "language must be a nonempty lowercase English key",
        )
    return None


def _sheet_issue(sheet: Any) -> dict[str, Any] | None:
    if not isinstance(sheet, str) or sheet not in KNOWLEDGE_SHEETS:
        return _issue(None, "knowledge_sheet", f"sheet must be one of {', '.join(KNOWLEDGE_SHEETS)}")
    return None


def _is_blank(value: Any) -> bool:
    return value is None or value == ""


def _is_nonempty_text(value: Any) -> bool:
    return isinstance(value, str) and bool(value.strip())


def _valid_integer(value: Any, minimum: int) -> bool:
    return isinstance(value, int) and not isinstance(value, bool) and value >= minimum


def _absolute(path: str | os.PathLike[str]) -> Path:
    return Path(path).expanduser().resolve()


def _header_map(worksheet: Any) -> dict[str, int]:
    result: dict[str, int] = {}
    for column, cell in enumerate(worksheet[1], start=1):
        if isinstance(cell.value, str) and cell.value not in result:
            result[cell.value] = column
    return result


def _row_is_blank(worksheet: Any, row: int, max_column: int) -> bool:
    return all(_is_blank(worksheet.cell(row, column).value) for column in range(1, max_column + 1))


def _row_record(worksheet: Any, row: int, headers: Mapping[str, int] | None = None) -> dict[str, Any]:
    if headers is None:
        headers = _header_map(worksheet)
    return {field: worksheet.cell(row, column).value for field, column in headers.items()}


def _load(path: Path) -> Any:
    return load_workbook(path, keep_vba=path.suffix.lower() == ".xlsm")


def _border_row(worksheet: Any, row: int, columns: Iterable[int]) -> bool:
    changed = False
    for column in columns:
        cell = worksheet.cell(row, column)
        if cell.border != _CELL_BORDER:
            cell.border = _CELL_BORDER
            changed = True
    return changed


def _set_mastery_dropdown(worksheet: Any, column: int) -> bool:
    """Replace rules on mastery data cells, preserving rules on all other cells."""

    letter = get_column_letter(column)
    expected = DataValidation(
        type="list", formula1=f'"{",".join(_MASTERY_OPTIONS)}"', allow_blank=True,
        showDropDown=False, showErrorMessage=True, errorStyle="stop",
        errorTitle="Invalid mastery", error="Choose a listed value or leave blank.",
    )
    expected.add(f"{letter}2:{letter}1048576")
    overlapping = [
        rule for rule in worksheet.data_validations.dataValidation
        if any(area.min_col <= column <= area.max_col and area.max_row >= 2 for area in rule.sqref.ranges)
    ]
    if overlapping == [expected]:
        return False
    for rule in overlapping:
        remaining = []
        for area in rule.sqref.ranges:
            if not area.min_col <= column <= area.max_col or area.max_row < 2:
                remaining.append(area)
                continue
            if area.min_row == 1:
                remaining.append(CellRange(min_col=area.min_col, max_col=area.max_col, min_row=1, max_row=1))
            for left, right in ((area.min_col, column - 1), (column + 1, area.max_col)):
                if left <= right:
                    remaining.append(CellRange(min_col=left, max_col=right, min_row=max(2, area.min_row), max_row=area.max_row))
        rule.sqref = MultiCellRange(remaining)
        if not remaining:
            worksheet.data_validations.dataValidation.remove(rule)
    worksheet.add_data_validation(expected)
    return True


def _apply_workbook_format(
    workbook: Any, *, hide_id: bool, header_fill: bool, borders: bool, mastery_dropdown: bool = False
) -> bool:
    """Apply selected defaults only to the required library fields."""

    changed = False
    for name, fields in HEADERS.items():
        worksheet = workbook[name]
        headers = _header_map(worksheet)
        columns = tuple(headers[field] for field in fields)
        if mastery_dropdown and name in KNOWLEDGE_SHEETS:
            changed |= _set_mastery_dropdown(worksheet, headers["mastery"])
        if hide_id and name in KNOWLEDGE_SHEETS:
            dimension = worksheet.column_dimensions[get_column_letter(headers["id"])]
            if not dimension.hidden:
                dimension.hidden = True
                changed = True
        if header_fill:
            for column in columns:
                cell = worksheet.cell(1, column)
                if cell.fill != _HEADER_FILL:
                    cell.fill = _HEADER_FILL
                    changed = True
        if borders:
            rows = (1, 2) if name == "Library" else (
                row for row in range(1, worksheet.max_row + 1)
                if row == 1 or not _is_blank(worksheet.cell(row, headers["item"]).value)
            )
            for row in rows:
                changed |= _border_row(worksheet, row, columns)
    return changed


class _FileWriteError(OSError):
    def __init__(self, detail: dict[str, Any]) -> None:
        super().__init__(detail["error"])
        self.detail = detail


def _file_error_detail(exc: Exception, path: Path, stage: str) -> dict[str, Any]:
    if isinstance(exc, _FileWriteError):
        return dict(exc.detail)
    return {
        "path": str(path),
        "directory": str(path.parent),
        "stage": stage,
        "error": str(exc),
        "error_type": type(exc).__name__,
        "errno": getattr(exc, "errno", None),
        "winerror": getattr(exc, "winerror", None),
    }


def _create_temporary(path: Path) -> Path:
    """Create a sibling file exclusively; retry name collisions, never permission errors."""

    for _ in range(5):
        temporary = path.with_name(f".{path.name}.{uuid.uuid4().hex}{path.suffix or '.tmp'}")
        try:
            descriptor = os.open(temporary, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
        except FileExistsError:
            continue
        os.close(descriptor)
        return temporary
    raise FileExistsError("temporary file names collided five times")


def _write_file_atomic(path: Path, writer: Callable[[Path], Any], *, replace_existing: bool = True) -> None:
    """Write a complete sibling file, then replace the destination without retrying failures."""

    temporary = None
    placeholder_created = False
    stage = "create_directory"
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        stage = "create_temporary"
        temporary = _create_temporary(path)
        stage = "write_temporary"
        writer(temporary)
        if not replace_existing:
            stage = "create_destination"
            descriptor = os.open(path, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
            placeholder_created = True
            os.close(descriptor)
        stage = "replace_destination"
        os.replace(temporary, path)
    except Exception as exc:
        detail = _file_error_detail(exc, path, stage)
        cleanup_errors = []
        for owned_path in ([path] if placeholder_created else []) + ([temporary] if temporary else []):
            try:
                owned_path.unlink(missing_ok=True)
            except OSError as cleanup_exc:
                cleanup_errors.append(_file_error_detail(cleanup_exc, owned_path, "remove_temporary"))
        if cleanup_errors:
            detail["cleanup_errors"] = cleanup_errors
        raise _FileWriteError(detail) from exc


def _save_workbook_atomic(workbook: Any, path: Path, *, replace_existing: bool = True) -> None:
    """Save beside the destination, then atomically publish the completed file."""

    _write_file_atomic(path, workbook.save, replace_existing=replace_existing)


def _read_registry(registry_path: Path) -> tuple[dict[str, str] | None, dict[str, Any] | None]:
    if not registry_path.exists():
        return {}, None
    try:
        data = json.loads(registry_path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        return None, _fail("registry_read_failed", _file_error_detail(exc, registry_path, "read_registry"))
    if not isinstance(data, dict) or not all(isinstance(key, str) and isinstance(value, str) for key, value in data.items()):
        return None, _fail(
            "registry_invalid",
            {"path": str(registry_path), "reason": "registry must be a JSON object mapping language keys to paths"},
        )
    return data, None


def _write_registry(registry_path: Path, bindings: Mapping[str, str]) -> None:
    content = json.dumps(dict(sorted(bindings.items())), ensure_ascii=False, indent=2) + "\n"
    _write_file_atomic(registry_path, lambda temporary: temporary.write_text(content, encoding="utf-8"))


def _validate_loaded(workbook: Any, path: Path, language: str) -> dict[str, Any]:
    errors: list[dict[str, Any]] = []
    header_maps: dict[str, dict[str, int]] = {}

    for sheet in SHEETS:
        if sheet not in workbook.sheetnames:
            errors.append(_issue(sheet, "required_sheet", f"required sheet {sheet!r} is missing"))
            continue
        headers = _header_map(workbook[sheet])
        header_maps[sheet] = headers
        for header in HEADERS[sheet]:
            if header not in headers:
                errors.append(_issue(sheet, "required_header", f"required header {header!r} is missing", row=1, header=header))

    if "Library" in header_maps and "language" in header_maps["Library"]:
        worksheet = workbook["Library"]
        column = header_maps["Library"]["language"]
        values = [
            (row, worksheet.cell(row, column).value)
            for row in range(2, worksheet.max_row + 1)
            if not _is_blank(worksheet.cell(row, column).value)
        ]
        if len(values) != 1:
            errors.append(
                _issue(
                    "Library",
                    "language_value_count",
                    "Library!language must contain exactly one nonblank data value",
                    rows=[row for row, _ in values],
                )
            )
        else:
            row, value = values[0]
            if row != 2:
                errors.append(_issue("Library", "language_row", "Library!language must be stored in row 2", row=row))
            if value != language:
                errors.append(
                    _issue(
                        "Library",
                        "language_match",
                        f"Library!language must equal expected key {language!r}",
                        row=row,
                        actual=value,
                        expected=language,
                    )
                )

    identifiers: dict[str, list[tuple[str, int]]] = {}
    identities: dict[tuple[str, str], list[int]] = {}
    for sheet in KNOWLEDGE_SHEETS:
        if sheet not in header_maps:
            continue
        headers = header_maps[sheet]
        worksheet = workbook[sheet]
        max_column = worksheet.max_column
        for row in range(2, worksheet.max_row + 1):
            if _row_is_blank(worksheet, row, max_column):
                continue

            for field in REQUIRED_TEXT[sheet]:
                if field in headers:
                    value = worksheet.cell(row, headers[field]).value
                    if not _is_nonempty_text(value):
                        errors.append(
                            _issue(sheet, "required_nonempty_text", f"{field} must be nonempty text", row=row, field=field)
                        )
            for field in OPTIONAL_TEXT[sheet]:
                if field in headers:
                    value = worksheet.cell(row, headers[field]).value
                    if not _is_blank(value) and not isinstance(value, str):
                        errors.append(_issue(sheet, "optional_text", f"{field} must be blank or text", row=row, field=field))

            if "frequency" in headers:
                value = worksheet.cell(row, headers["frequency"]).value
                if not _valid_integer(value, 1):
                    errors.append(
                        _issue(
                            sheet,
                            "frequency_integer_minimum",
                            "frequency must be a non-Boolean integer greater than or equal to 1",
                            row=row,
                            actual=value,
                        )
                    )
            if "mastery" in headers:
                value = worksheet.cell(row, headers["mastery"]).value
                if not _is_blank(value) and value not in MASTERY_VALUES:
                    errors.append(
                        _issue(
                            sheet,
                            "mastery_enum",
                            "mastery must be blank, unmastered, unfamiliar, or mastered",
                            row=row,
                            actual=value,
                        )
                    )
            if "example" in headers and "example_translation" in headers:
                example_blank = _is_blank(worksheet.cell(row, headers["example"]).value)
                translation_blank = _is_blank(worksheet.cell(row, headers["example_translation"]).value)
                if example_blank != translation_blank:
                    errors.append(
                        _issue(
                            sheet,
                            "example_pair",
                            "example and example_translation must both be blank or both be nonblank",
                            row=row,
                        )
                    )

            if "id" in headers:
                identifier = worksheet.cell(row, headers["id"]).value
                if _is_nonempty_text(identifier):
                    identifiers.setdefault(identifier, []).append((sheet, row))
            if "item" in headers:
                item = worksheet.cell(row, headers["item"]).value
                if _is_nonempty_text(item):
                    identities.setdefault((sheet, normalize_item(item)), []).append(row)

    for identifier, locations in identifiers.items():
        if len(locations) > 1:
            errors.append(
                _issue(
                    locations[0][0],
                    "unique_id",
                    f"id {identifier!r} must be unique across the workbook",
                    row=locations[0][1],
                    locations=[{"sheet": sheet, "row": row} for sheet, row in locations],
                )
            )
    for (sheet, item), rows in identities.items():
        if len(rows) > 1:
            errors.append(
                _issue(
                    sheet,
                    "unique_normalized_identity",
                    f"normalized item identity {item!r} is duplicated in {sheet}",
                    row=rows[0],
                    rows=rows,
                )
            )

    if errors:
        return _fail("validation_failed", {"errors": errors}, valid=False, path=str(path), language=language)
    return _ok(valid=True, path=str(path), language=language, errors=[])


def validate_workbook(workbook_path: str | os.PathLike[str], language: str) -> dict[str, Any]:
    """Validate a workbook against the exact schema and expected language."""

    issue = _language_issue(language)
    if issue:
        return _fail("invalid_language", issue, valid=False)
    try:
        path = _absolute(workbook_path)
    except (TypeError, OSError, ValueError) as exc:
        return _fail("invalid_path", {"error": str(exc)}, valid=False)
    if not path.is_file():
        error = _issue(None, "workbook_exists", "workbook file does not exist")
        return _fail("workbook_missing", {"errors": [error]}, valid=False, path=str(path), language=language)
    try:
        workbook = _load(path)
    except Exception as exc:
        error = _issue(None, "workbook_readable", f"workbook could not be opened: {exc}")
        return _fail("workbook_open_failed", {"errors": [error]}, valid=False, path=str(path), language=language)
    try:
        return _validate_loaded(workbook, path, language)
    finally:
        workbook.close()


def _check_workbook_writable(path: Path) -> dict[str, Any]:
    """Check the existing file and sibling creation; replacement can still fail later."""

    stage = "open_existing"
    try:
        with path.open("r+b"):
            pass
        stage = "create_temporary"
        temporary = _create_temporary(path)
        stage = "remove_temporary"
        temporary.unlink()
    except OSError as exc:
        return _fail("workbook_not_writable", _file_error_detail(exc, path, stage), path=str(path))
    return _ok(path=str(path))


def bind_workbook(
    language: str,
    workbook_path: str | os.PathLike[str],
    registry_path: str | os.PathLike[str] = DEFAULT_REGISTRY_PATH,
) -> dict[str, Any]:
    """Bind a caller-selected workbook; this function never searches or chooses candidates."""

    issue = _language_issue(language)
    if issue:
        return _fail("invalid_language", issue, language=language)
    try:
        path = _absolute(workbook_path)
        registry = _absolute(registry_path)
    except (TypeError, OSError, ValueError) as exc:
        return _fail("invalid_path", {"error": str(exc)}, language=language)

    validation = validate_workbook(path, language)
    if not validation["ok"]:
        return _fail("workbook_invalid", validation["detail"], language=language, path=str(path))
    writable = _check_workbook_writable(path)
    if not writable["ok"]:
        return {**writable, "language": language}

    bindings, failure = _read_registry(registry)
    if failure:
        return failure
    assert bindings is not None
    previous = bindings.get(language)
    bindings[language] = str(path)
    try:
        _write_registry(registry, bindings)
    except OSError as exc:
        return _fail("registry_write_failed", _file_error_detail(exc, registry, "write_registry"), language=language)
    return _ok(
        action="rebound" if previous and previous != str(path) else "bound",
        language=language,
        path=str(path),
        previous_path=previous,
        registry_path=str(registry),
    )


def resolve_workbook(
    language: str, registry_path: str | os.PathLike[str] = DEFAULT_REGISTRY_PATH
) -> dict[str, Any]:
    """Resolve only the current binding; never search for or select a replacement."""

    issue = _language_issue(language)
    if issue:
        return _fail("invalid_language", issue, language=language)
    try:
        registry = _absolute(registry_path)
    except (TypeError, OSError, ValueError) as exc:
        return _fail("invalid_path", {"error": str(exc)}, language=language)
    bindings, failure = _read_registry(registry)
    if failure:
        return failure
    assert bindings is not None
    stored = bindings.get(language)
    if stored is None:
        return _fail("binding_not_found", {"language": language, "registry_path": str(registry)}, language=language)
    path = Path(stored)
    if not path.is_absolute():
        return _fail("binding_invalid", {"language": language, "path": stored, "reason": "bound path is not absolute"})
    if not path.is_file():
        return _fail("bound_workbook_missing", {"language": language, "path": str(path)}, language=language, path=str(path))
    return _ok(language=language, path=str(path), registry_path=str(registry))


def search_workbook_candidates(
    language: str,
    directory: str | os.PathLike[str],
    filename: str,
) -> dict[str, Any]:
    """Recursively discover and validate candidates without selecting or binding one."""

    issue = _language_issue(language)
    if issue:
        return _fail("invalid_language", issue, language=language)
    try:
        root = _absolute(directory)
    except (TypeError, OSError, ValueError) as exc:
        return _fail("invalid_path", {"error": str(exc)}, language=language)
    if not root.is_dir():
        return _fail(
            "search_directory_missing",
            {"directory": str(root), "reason": "search directory does not exist or is not a directory"},
            language=language,
        )
    if not isinstance(filename, str) or not filename or Path(filename).name != filename:
        return _fail(
            "invalid_filename",
            {"filename": filename, "reason": "filename must be one file name without a directory path"},
            language=language,
            directory=str(root),
        )

    def raise_walk_error(error: OSError) -> None:
        raise error

    try:
        matches = sorted(
            Path(current_root) / name
            for current_root, _, files in os.walk(root, onerror=raise_walk_error)
            for name in files
            if name == filename
        )
    except OSError as exc:
        return _fail(
            "candidate_search_failed",
            {"directory": str(root), "filename": filename, "error": str(exc)},
            language=language,
        )

    valid_candidates: list[dict[str, Any]] = []
    invalid_matches: list[dict[str, Any]] = []
    for path in matches:
        validation = validate_workbook(path, language)
        if not validation["ok"]:
            invalid_matches.append(
                {
                    "ok": False,
                    "path": str(path),
                    "reason": "workbook_invalid",
                    "detail": validation["detail"],
                }
            )
            continue
        writable = _check_workbook_writable(path)
        if not writable["ok"]:
            invalid_matches.append(
                {
                    "ok": False,
                    "path": str(path),
                    "reason": writable["reason"],
                    "detail": writable["detail"],
                }
            )
            continue
        valid_candidates.append({"ok": True, "path": str(path), "validation": validation})

    if not matches:
        outcome = "no_matches"
    elif not valid_candidates:
        outcome = "matches_all_invalid"
    elif len(valid_candidates) == 1:
        outcome = "one_valid_candidate"
    else:
        outcome = "multiple_valid_candidates"
    return _ok(
        action="search",
        outcome=outcome,
        language=language,
        directory=str(root),
        filename=filename,
        match_count=len(matches),
        valid_count=len(valid_candidates),
        invalid_count=len(invalid_matches),
        valid_candidates=valid_candidates,
        invalid_matches=invalid_matches,
    )


def _rollback_workbook_move(destination: Path, source: Path, language: str) -> dict[str, Any]:
    try:
        shutil.move(str(destination), str(source))
    except Exception as exc:
        return _fail(
            "workbook_move_rollback_failed",
            {"source_path": str(source), "destination_path": str(destination), "error": str(exc)},
        )
    validation = validate_workbook(source, language)
    if not validation["ok"]:
        return _fail(
            "workbook_move_rollback_failed",
            {"source_path": str(source), "destination_path": str(destination), "validation": validation},
        )
    return _ok(source_path=str(source), destination_path=str(destination))


def move_workbook(
    language: str,
    destination_path: str | os.PathLike[str],
    registry_path: str | os.PathLike[str] = DEFAULT_REGISTRY_PATH,
) -> dict[str, Any]:
    """Move or rename the bound workbook to an explicit path, then validate and rebind it."""

    issue = _language_issue(language)
    if issue:
        return _fail("invalid_language", issue, language=language)
    resolved = resolve_workbook(language, registry_path)
    if not resolved["ok"]:
        return resolved
    source = Path(resolved["path"])
    try:
        destination = _absolute(destination_path)
    except (TypeError, OSError, ValueError) as exc:
        return _fail("invalid_path", {"error": str(exc)}, language=language, source_path=str(source))
    if destination.exists():
        return _fail(
            "destination_exists",
            {"source_path": str(source), "destination_path": str(destination), "reason": "destination will not be overwritten"},
            language=language,
        )
    if not destination.parent.is_dir():
        return _fail(
            "destination_parent_missing",
            {"source_path": str(source), "destination_path": str(destination)},
            language=language,
        )

    source_validation = validate_workbook(source, language)
    if not source_validation["ok"]:
        return _fail(
            "workbook_invalid",
            source_validation["detail"],
            language=language,
            source_path=str(source),
            destination_path=str(destination),
        )
    try:
        shutil.move(str(source), str(destination))
    except Exception as exc:
        cleanup_error = None
        if destination.exists():
            if source.exists():
                try:
                    destination.unlink()
                except OSError as cleanup_exc:
                    cleanup_error = str(cleanup_exc)
            else:
                rollback = _rollback_workbook_move(destination, source, language)
                if not rollback["ok"]:
                    cleanup_error = rollback
        return _fail(
            "workbook_move_failed",
            {
                "source_path": str(source),
                "destination_path": str(destination),
                "error": str(exc),
                "rollback_error": cleanup_error,
            },
            language=language,
        )

    destination_validation = validate_workbook(destination, language)
    if not destination_validation["ok"]:
        rollback = _rollback_workbook_move(destination, source, language)
        if not rollback["ok"]:
            return {**rollback, "language": language, "original_failure": destination_validation}
        return _fail(
            "destination_validation_failed",
            {
                "source_path": str(source),
                "destination_path": str(destination),
                "validation": destination_validation,
                "rolled_back": True,
            },
            language=language,
        )

    binding = bind_workbook(language, destination, registry_path)
    if not binding["ok"]:
        rollback = _rollback_workbook_move(destination, source, language)
        if not rollback["ok"]:
            return {**rollback, "language": language, "original_failure": binding}
        return _fail(
            "binding_update_failed",
            {
                "source_path": str(source),
                "destination_path": str(destination),
                "binding": binding,
                "rolled_back": True,
            },
            language=language,
        )
    return _ok(
        action="renamed" if source.parent == destination.parent else "moved",
        language=language,
        source_path=str(source),
        path=str(destination),
        registry_path=binding["registry_path"],
    )


def create_workbook(
    language: str,
    workbook_path: str | os.PathLike[str],
    registry_path: str | os.PathLike[str] = DEFAULT_REGISTRY_PATH,
    *,
    bind: bool = True,
) -> dict[str, Any]:
    """Create a schema-valid workbook with default display, without overwriting an existing file."""

    issue = _language_issue(language)
    if issue:
        return _fail("invalid_language", issue, language=language)
    try:
        path = _absolute(workbook_path)
    except (TypeError, OSError, ValueError) as exc:
        return _fail("invalid_path", {"error": str(exc)}, language=language)
    if path.exists():
        return _fail("workbook_exists", {"path": str(path), "reason": "existing workbooks are never overwritten"}, language=language)

    workbook = Workbook()
    library_sheet = workbook.active
    library_sheet.title = "Library"
    for sheet in KNOWLEDGE_SHEETS:
        workbook.create_sheet(sheet)
    for sheet, headers in HEADERS.items():
        workbook[sheet].append(headers)
    library_sheet["A2"] = language
    try:
        _apply_workbook_format(workbook, hide_id=True, header_fill=True, borders=True, mastery_dropdown=True)
        _save_workbook_atomic(workbook, path, replace_existing=False)
    except Exception as exc:
        return _fail("workbook_save_failed", _file_error_detail(exc, path, "save_workbook"), language=language)
    finally:
        workbook.close()

    validation = validate_workbook(path, language)
    if not validation["ok"]:
        return _fail("created_workbook_invalid", validation["detail"], language=language, path=str(path))
    if bind:
        binding = bind_workbook(language, path, registry_path)
        if not binding["ok"]:
            return _fail(
                "binding_failed",
                {"path": str(path), "binding": binding, "workbook_created": True},
                language=language,
                path=str(path),
            )
    return _ok(action="created", language=language, path=str(path), bound=bind)


def initialize_workbook(
    language: str,
    folder: str | os.PathLike[str],
    registry_path: str | os.PathLike[str] = DEFAULT_REGISTRY_PATH,
) -> dict[str, Any]:
    """Create and bind the schema's default filename inside ``folder``."""

    issue = _language_issue(language)
    if issue:
        return _fail("invalid_language", issue, language=language)
    try:
        path = Path(folder) / f"{language}_learning_library.xlsx"
    except (TypeError, OSError, ValueError) as exc:
        return _fail("invalid_path", {"error": str(exc)}, language=language)
    return create_workbook(language, path, registry_path, bind=True)


def _open_bound(language: str, registry_path: str | os.PathLike[str]) -> tuple[Any, Path] | dict[str, Any]:
    resolved = resolve_workbook(language, registry_path)
    if not resolved["ok"]:
        return resolved
    path = Path(resolved["path"])
    try:
        workbook = _load(path)
    except Exception as exc:
        return _fail("workbook_open_failed", {"path": str(path), "error": str(exc)}, language=language)
    validation = _validate_loaded(workbook, path, language)
    if not validation["ok"]:
        workbook.close()
        return _fail("workbook_invalid", validation["detail"], language=language, path=str(path))
    return workbook, path


def format_workbook(
    language: str,
    registry_path: str | os.PathLike[str] = DEFAULT_REGISTRY_PATH,
    *,
    hide_id: bool = False,
    header_fill: bool = False,
    borders: bool = False,
    mastery_dropdown: bool = False,
) -> dict[str, Any]:
    """Explicitly apply selected display defaults to a validated bound workbook.

    False options preserve the corresponding display. Never called by binding,
    lookup, or record updates. Record values and registry bindings remain unchanged.
    """

    options = {"hide_id": hide_id, "header_fill": header_fill, "borders": borders, "mastery_dropdown": mastery_dropdown}
    invalid = [name for name, value in options.items() if not isinstance(value, bool)]
    if invalid:
        return _fail(
            "format_options_invalid",
            _issue(None, "format_boolean_options", "format options must be booleans", fields=invalid),
            persisted=False,
        )
    opened = _open_bound(language, registry_path)
    if isinstance(opened, dict):
        return {**opened, "persisted": False}
    workbook, path = opened
    stage = "format_workbook"
    try:
        changed = _apply_workbook_format(workbook, **options)
        if changed:
            stage = "save_workbook"
            _save_workbook_atomic(workbook, path)
        return _ok(action="formatted" if changed else "unchanged", language=language, path=str(path), persisted=changed)
    except Exception as exc:
        return _fail(
            "workbook_save_failed" if stage == "save_workbook" else "workbook_format_failed",
            _file_error_detail(exc, path, stage),
            language=language, path=str(path), persisted=False,
        )
    finally:
        workbook.close()


def _find_row(workbook: Any, sheet: str, item: str) -> int | None:
    worksheet = workbook[sheet]
    column = _header_map(worksheet)["item"]
    for row in range(2, worksheet.max_row + 1):
        value = worksheet.cell(row, column).value
        if isinstance(value, str) and normalize_item(value) == item:
            return row
    return None


def lookup_record(
    language: str,
    sheet: str,
    item: str,
    registry_path: str | os.PathLike[str] = DEFAULT_REGISTRY_PATH,
) -> dict[str, Any]:
    """Look up one normalized identity in the language's bound workbook."""

    language_error = _language_issue(language)
    if language_error:
        return _fail("invalid_language", language_error, language=language)
    sheet_error = _sheet_issue(sheet)
    if sheet_error:
        return _fail("invalid_sheet", sheet_error, language=language, sheet=sheet)
    try:
        normalized = normalize_item(item)
    except TypeError as exc:
        return _fail("invalid_item", _issue(sheet, "item_text", str(exc)), language=language, sheet=sheet)
    if not normalized:
        return _fail("invalid_item", _issue(sheet, "item_nonempty", "item must not be blank"), language=language, sheet=sheet)

    opened = _open_bound(language, registry_path)
    if isinstance(opened, dict):
        return opened
    workbook, path = opened
    try:
        row = _find_row(workbook, sheet, normalized)
        if row is None:
            return _ok(language=language, path=str(path), sheet=sheet, item=normalized, exists=False, row=None, record=None)
        return _ok(
            language=language,
            path=str(path),
            sheet=sheet,
            item=normalized,
            exists=True,
            row=row,
            record=_row_record(workbook[sheet], row),
        )
    finally:
        workbook.close()


def read_workbook(
    language: str, registry_path: str | os.PathLike[str] = DEFAULT_REGISTRY_PATH
) -> dict[str, Any]:
    """Read every nonblank committed record from a validated bound workbook."""

    opened = _open_bound(language, registry_path)
    if isinstance(opened, dict):
        return opened
    workbook, path = opened
    try:
        records: dict[str, list[dict[str, Any]]] = {}
        for sheet in KNOWLEDGE_SHEETS:
            worksheet = workbook[sheet]
            headers = _header_map(worksheet)
            max_column = worksheet.max_column
            records[sheet] = [
                {"row": row, "record": _row_record(worksheet, row, headers)}
                for row in range(2, worksheet.max_row + 1)
                if not _row_is_blank(worksheet, row, max_column)
            ]
        return _ok(language=language, path=str(path), records=records)
    finally:
        workbook.close()


def _parse_collocations(value: Any, *, allow_blank: bool) -> tuple[list[str] | None, str | None]:
    if _is_blank(value):
        return ([], None) if allow_blank else (None, "collocations must contain nonempty text")
    if isinstance(value, str):
        raw_values: Iterable[Any] = value.split(";")
    elif isinstance(value, Sequence) and not isinstance(value, (str, bytes, bytearray)):
        raw_values = value
    else:
        return None, "collocations must be text or a list of text values"

    parsed: list[str] = []
    seen: set[str] = set()
    for entry in raw_values:
        if not isinstance(entry, str):
            return None, "every collocation must be text"
        normalized = normalize_item(entry)
        if not normalized:
            if allow_blank:
                continue
            return None, "collocations must not contain blank entries"
        if normalized not in seen:
            seen.add(normalized)
            parsed.append(normalized)
    if not parsed and not allow_blank:
        return None, "collocations must contain nonempty text"
    return parsed, None


def _prepare_insert(sheet: str, record: Any) -> tuple[dict[str, Any] | None, dict[str, Any] | None]:
    if not isinstance(record, Mapping):
        return None, _issue(sheet, "record_object", "record must be an object")
    allowed = set(HEADERS[sheet]) - {"id"}
    unexpected = sorted(set(record) - allowed, key=str)
    if unexpected:
        return None, _issue(sheet, "insert_allowed_fields", "insert contains fields that are not allowed", fields=unexpected)

    prepared: dict[str, Any] = {field: None for field in allowed}
    try:
        prepared["item"] = normalize_item(record.get("item"))
    except TypeError as exc:
        return None, _issue(sheet, "item_text", str(exc), field="item")
    if not prepared["item"]:
        return None, _issue(sheet, "item_nonempty", "item must not be blank", field="item")

    for field in REQUIRED_TEXT[sheet]:
        if field in {"id", "item"}:
            continue
        value = record.get(field)
        if not _is_nonempty_text(value):
            return None, _issue(sheet, "required_nonempty_text", f"{field} must be nonempty text", field=field)
        prepared[field] = value

    if "frequency" not in record or not _valid_integer(record.get("frequency"), 1):
        return None, _issue(
            sheet,
            "frequency_integer_minimum",
            "insert frequency must be a caller-supplied non-Boolean integer greater than or equal to 1",
            field="frequency",
            actual=record.get("frequency"),
        )
    prepared["frequency"] = record["frequency"]

    for field in OPTIONAL_TEXT[sheet]:
        if field == "collocations":
            continue
        value = record.get(field)
        if not _is_blank(value) and not isinstance(value, str):
            return None, _issue(sheet, "optional_text", f"{field} must be blank or text", field=field)
        prepared[field] = None if _is_blank(value) else value

    if sheet == "Vocabulary":
        collocations, error = _parse_collocations(record.get("collocations"), allow_blank=True)
        if error:
            return None, _issue(sheet, "collocations_format", error, field="collocations")
        prepared["collocations"] = "; ".join(collocations or []) or None

    example_blank = _is_blank(prepared.get("example"))
    translation_blank = _is_blank(prepared.get("example_translation"))
    if example_blank != translation_blank:
        return None, _issue(
            sheet,
            "example_pair",
            "example and example_translation must both be blank or both be nonblank",
        )

    mastery = record.get("mastery")
    if _is_blank(mastery):
        prepared["mastery"] = None
    elif not isinstance(mastery, str) or mastery not in MASTERY_VALUES:
        return None, _issue(
            sheet,
            "mastery_enum",
            "mastery must be blank, unmastered, unfamiliar, or mastered",
            field="mastery",
            actual=mastery,
        )
    else:
        prepared["mastery"] = mastery
    return prepared, None


def _apply_insert(workbook: Any, sheet: str, record: Any) -> tuple[dict[str, Any], bool]:
    prepared, error = _prepare_insert(sheet, record)
    if error:
        return _fail("record_invalid", error, mode="insert", sheet=sheet), False
    assert prepared is not None
    existing_row = _find_row(workbook, sheet, prepared["item"])
    if existing_row is not None:
        return _fail(
            "identity_exists",
            _issue(sheet, "insert_requires_absent_identity", "insert is allowed only when normalized identity is absent", row=existing_row),
            mode="insert",
            sheet=sheet,
            item=prepared["item"],
        ), False

    worksheet = workbook[sheet]
    headers = _header_map(worksheet)
    row = worksheet.max_row + 1
    prepared["id"] = str(uuid.uuid4())
    for field in HEADERS[sheet]:
        worksheet.cell(row, headers[field]).value = prepared.get(field)
    _border_row(worksheet, row, (headers[field] for field in HEADERS[sheet]))
    return (
        _ok(mode="insert", status="inserted", sheet=sheet, item=prepared["item"], row=row, record=_row_record(worksheet, row)),
        True,
    )


def _prepare_update(sheet: str, updates: Any) -> tuple[dict[str, Any] | None, dict[str, Any] | None]:
    if updates is None:
        return {}, None
    if not isinstance(updates, Mapping):
        return None, _issue(sheet, "updates_object", "updates must be an object")
    unexpected = sorted(set(updates) - UPDATE_FIELDS[sheet], key=str)
    if unexpected:
        return None, _issue(sheet, "commit_update_allowed_fields", "commit_update contains fields that are not allowed", fields=unexpected)

    prepared: dict[str, Any] = {}
    for field, value in updates.items():
        if field == "collocations":
            collocations, error = _parse_collocations(value, allow_blank=False)
            if error:
                return None, _issue(sheet, "collocations_format", error, field=field)
            prepared[field] = collocations
        elif field == "mastery":
            if _is_blank(value):
                prepared[field] = None
            elif not isinstance(value, str) or value not in MASTERY_VALUES:
                return None, _issue(
                    sheet,
                    "mastery_enum",
                    "an explicit mastery update must be blank, unmastered, unfamiliar, or mastered",
                    field=field,
                    actual=value,
                )
            else:
                prepared[field] = value
        elif field in OPTIONAL_TEXT[sheet]:
            if not _is_blank(value) and not isinstance(value, str):
                return None, _issue(sheet, "optional_text", f"{field} must be blank or text", field=field)
            prepared[field] = None if _is_blank(value) else value
        else:
            if not _is_nonempty_text(value):
                return None, _issue(sheet, "required_nonempty_text", f"{field} must be nonempty text", field=field)
            prepared[field] = value
    return prepared, None


def _apply_commit_update(
    workbook: Any, sheet: str, item: Any, delta: Any, updates: Any
) -> tuple[dict[str, Any], bool]:
    try:
        normalized = normalize_item(item)
    except TypeError as exc:
        return _fail("invalid_item", _issue(sheet, "item_text", str(exc)), mode="commit_update", sheet=sheet), False
    if not normalized:
        return _fail(
            "invalid_item", _issue(sheet, "item_nonempty", "item must not be blank"), mode="commit_update", sheet=sheet
        ), False
    if not _valid_integer(delta, 0):
        return _fail(
            "invalid_delta",
            _issue(sheet, "nonnegative_integer_delta", "delta must be a non-Boolean integer greater than or equal to 0"),
            mode="commit_update",
            sheet=sheet,
            item=normalized,
        ), False
    prepared, error = _prepare_update(sheet, updates)
    if error:
        return _fail("updates_invalid", error, mode="commit_update", sheet=sheet, item=normalized), False
    assert prepared is not None

    row = _find_row(workbook, sheet, normalized)
    if row is None:
        return _fail(
            "identity_not_found",
            _issue(sheet, "commit_update_requires_existing_identity", "commit_update is allowed only for an existing identity"),
            mode="commit_update",
            sheet=sheet,
            item=normalized,
        ), False

    worksheet = workbook[sheet]
    headers = _header_map(worksheet)
    values: dict[str, Any] = {}
    for field, value in prepared.items():
        if field in {"collocations", "mastery"}:
            continue
        values[field] = value

    if "collocations" in prepared:
        existing, error = _parse_collocations(worksheet.cell(row, headers["collocations"]).value, allow_blank=True)
        assert error is None
        merged = list(existing or [])
        seen = set(merged)
        for collocation in prepared["collocations"]:
            if collocation not in seen:
                seen.add(collocation)
                merged.append(collocation)
        values["collocations"] = "; ".join(merged) or None
    if "mastery" in prepared:
        values["mastery"] = prepared["mastery"]

    resulting_example = values.get("example", worksheet.cell(row, headers["example"]).value)
    resulting_translation = values.get(
        "example_translation", worksheet.cell(row, headers["example_translation"]).value
    )
    if _is_blank(resulting_example) != _is_blank(resulting_translation):
        return _fail(
            "updates_invalid",
            _issue(
                sheet,
                "example_pair",
                "example and example_translation must both be blank or both be nonblank after commit_update",
                row=row,
            ),
            mode="commit_update",
            sheet=sheet,
            item=normalized,
        ), False

    values["frequency"] = worksheet.cell(row, headers["frequency"]).value + delta
    changed = any(worksheet.cell(row, headers[field]).value != value for field, value in values.items())
    if changed:
        for field, value in values.items():
            worksheet.cell(row, headers[field]).value = value
    return (
        _ok(
            mode="commit_update",
            status="updated" if changed else "unchanged",
            sheet=sheet,
            item=normalized,
            row=row,
            record=_row_record(worksheet, row),
        ),
        changed,
    )


def _apply_passive_count(workbook: Any, sheet: str, item: Any, delta: Any) -> tuple[dict[str, Any], bool]:
    try:
        normalized = normalize_item(item)
    except TypeError as exc:
        return _fail("invalid_item", _issue(sheet, "item_text", str(exc)), mode="passive_count", sheet=sheet), False
    if not normalized:
        return _fail(
            "invalid_item", _issue(sheet, "item_nonempty", "item must not be blank"), mode="passive_count", sheet=sheet
        ), False
    if not _valid_integer(delta, 0):
        return _fail(
            "invalid_delta",
            _issue(sheet, "nonnegative_integer_delta", "delta must be a non-Boolean integer greater than or equal to 0"),
            mode="passive_count",
            sheet=sheet,
            item=normalized,
        ), False
    row = _find_row(workbook, sheet, normalized)
    if row is None:
        return _fail(
            "identity_not_found",
            _issue(sheet, "passive_count_requires_existing_identity", "passive_count is allowed only for an existing identity"),
            mode="passive_count",
            sheet=sheet,
            item=normalized,
        ), False
    worksheet = workbook[sheet]
    headers = _header_map(worksheet)
    frequency_cell = worksheet.cell(row, headers["frequency"])
    if delta:
        frequency_cell.value += delta
    return (
        _ok(
            mode="passive_count",
            status="updated" if delta else "unchanged",
            sheet=sheet,
            item=normalized,
            row=row,
            frequency=frequency_cell.value,
            record=_row_record(worksheet, row),
        ),
        bool(delta),
    )


def _validate_operation_shape(operation: Any) -> tuple[dict[str, Any] | None, dict[str, Any] | None]:
    if not isinstance(operation, Mapping):
        return None, _issue(None, "operation_object", "operation must be an object")
    mode = operation.get("mode")
    allowed = {
        "insert": {"mode", "language", "sheet", "record"},
        "commit_update": {"mode", "language", "sheet", "item", "delta", "updates"},
        "passive_count": {"mode", "language", "sheet", "item", "delta"},
    }
    if not isinstance(mode, str) or mode not in allowed:
        return None, _issue(None, "operation_mode", "mode must be insert, commit_update, or passive_count")
    unexpected = sorted(set(operation) - allowed[mode], key=str)
    if unexpected:
        rule = "passive_count_frequency_only" if mode == "passive_count" else f"{mode}_allowed_fields"
        reason = (
            "passive_count may contain only language, sheet, item, and delta"
            if mode == "passive_count"
            else f"{mode} contains fields that are not allowed"
        )
        return None, _issue(operation.get("sheet"), rule, reason, fields=unexpected)
    required = allowed[mode] - ({"updates", "delta"} if mode == "commit_update" else set())
    missing = sorted((field for field in required if field not in operation), key=str)
    if missing:
        return None, _issue(operation.get("sheet"), "required_operation_fields", "operation is missing required fields", fields=missing)
    shaped = dict(operation)
    if mode == "commit_update":
        shaped.setdefault("delta", 0)
        shaped.setdefault("updates", {})
    return shaped, None


def execute_operation(
    operation: Mapping[str, Any], registry_path: str | os.PathLike[str] = DEFAULT_REGISTRY_PATH
) -> dict[str, Any]:
    """Validate and execute exactly one insert, commit_update, or passive_count."""

    result = batch_execute_operations([operation], registry_path)["results"][0]
    return {key: value for key, value in result.items() if key != "index"}


def insert_record(
    language: str,
    sheet: str,
    record: Mapping[str, Any],
    registry_path: str | os.PathLike[str] = DEFAULT_REGISTRY_PATH,
) -> dict[str, Any]:
    return execute_operation({"mode": "insert", "language": language, "sheet": sheet, "record": record}, registry_path)


def commit_update(
    language: str,
    sheet: str,
    item: str,
    *,
    delta: int = 0,
    updates: Mapping[str, Any] | None = None,
    registry_path: str | os.PathLike[str] = DEFAULT_REGISTRY_PATH,
) -> dict[str, Any]:
    """Update an existing record; every field present in updates is an explicit caller-authorized edit."""

    return execute_operation(
        {
            "mode": "commit_update",
            "language": language,
            "sheet": sheet,
            "item": item,
            "delta": delta,
            "updates": {} if updates is None else updates,
        },
        registry_path,
    )


def passive_count(
    language: str,
    sheet: str,
    item: str,
    delta: int,
    registry_path: str | os.PathLike[str] = DEFAULT_REGISTRY_PATH,
    **forbidden_fields: Any,
) -> dict[str, Any]:
    operation = {"mode": "passive_count", "language": language, "sheet": sheet, "item": item, "delta": delta}
    operation.update(forbidden_fields)
    return execute_operation(operation, registry_path)


def _batch_operation(value: Any) -> tuple[dict[str, Any] | None, dict[str, Any] | None]:
    if isinstance(value, Mapping):
        operation = dict(value)
        if operation.get("mode", "passive_count") != "passive_count":
            return None, _issue(operation.get("sheet"), "passive_count_mode", "batch accepts passive_count operations only")
        operation.pop("mode", None)
    elif isinstance(value, Sequence) and not isinstance(value, (str, bytes, bytearray)) and len(value) == 4:
        operation = dict(zip(("language", "sheet", "item", "delta"), value))
    else:
        return None, _issue(None, "batch_record_shape", "each batch record must be an object or a four-value sequence")
    unexpected = sorted(set(operation) - {"language", "sheet", "item", "delta"}, key=str)
    if unexpected:
        return None, _issue(
            operation.get("sheet"),
            "passive_count_frequency_only",
            "passive_count may contain only language, sheet, item, and delta",
            fields=unexpected,
        )
    missing = sorted((field for field in ("language", "sheet", "item", "delta") if field not in operation), key=str)
    if missing:
        return None, _issue(operation.get("sheet"), "required_operation_fields", "operation is missing required fields", fields=missing)
    return operation, None


def batch_passive_count(
    operations: Sequence[Any], registry_path: str | os.PathLike[str] = DEFAULT_REGISTRY_PATH
) -> dict[str, Any]:
    """Apply independent passive counts, saving each affected workbook once."""

    return _execute_batch(operations, registry_path, passive_only=True)


def batch_execute_operations(
    operations: Sequence[Mapping[str, Any]], registry_path: str | os.PathLike[str] = DEFAULT_REGISTRY_PATH
) -> dict[str, Any]:
    """Execute explicit modes in input order, loading and saving each language's workbook once.

    Results keep input indices. Invalid records do not block valid records. A failed
    save fails every staged operation for that workbook; other languages continue.
    Returned records reflect the final row. No candidate selection or retries occur.
    """

    return _execute_batch(operations, registry_path)


def _apply_operation(workbook: Any, operation: Mapping[str, Any]) -> tuple[dict[str, Any], bool]:
    sheet = operation["sheet"]
    if operation["mode"] == "insert":
        return _apply_insert(workbook, sheet, operation["record"])
    if operation["mode"] == "commit_update":
        return _apply_commit_update(workbook, sheet, operation["item"], operation["delta"], operation["updates"])
    return _apply_passive_count(workbook, sheet, operation["item"], operation["delta"])


def _execute_batch(
    operations: Sequence[Any], registry_path: str | os.PathLike[str], *, passive_only: bool = False
) -> dict[str, Any]:
    if not isinstance(operations, Sequence) or isinstance(operations, (str, bytes, bytearray)):
        return _fail("batch_invalid", {"reason": "operations must be a list or tuple"}, results=[])
    results: list[dict[str, Any] | None] = [None] * len(operations)
    groups: dict[str, list[tuple[int, dict[str, Any]]]] = {}

    for index, value in enumerate(operations):
        if passive_only:
            operation, error = _batch_operation(value)
            if not error:
                value = {**operation, "mode": "passive_count"}
        else:
            error = None
        if not error:
            operation, error = _validate_operation_shape(value)
        if error:
            results[index] = _fail("operation_invalid", error, index=index)
            if passive_only:
                results[index]["mode"] = "passive_count"
            continue
        assert operation is not None
        language = operation["language"]
        sheet = operation["sheet"]
        language_error = _language_issue(language)
        sheet_error = _sheet_issue(sheet)
        if language_error or sheet_error:
            results[index] = _fail(
                "invalid_language" if language_error else "invalid_sheet",
                language_error or sheet_error,
                index=index,
                mode=operation["mode"],
                language=language,
                sheet=sheet,
            )
            continue
        groups.setdefault(language, []).append((index, operation))

    for language, grouped in groups.items():
        opened = _open_bound(language, registry_path)
        if isinstance(opened, dict):
            for index, operation in grouped:
                results[index] = {**opened, "index": index, "mode": operation["mode"]}
            continue
        workbook, path = opened
        try:
            staged: list[tuple[int, dict[str, Any], bool]] = []
            execution_failure = None
            for index, operation in grouped:
                try:
                    result, changed = _apply_operation(workbook, operation)
                except Exception as exc:
                    execution_failure = _file_error_detail(exc, path, "apply_operation")
                    break
                if not result["ok"]:
                    results[index] = {**result, "index": index, "language": language, "path": str(path)}
                else:
                    staged.append((index, result, changed))

            # An unexpected exception may have partially changed memory. Discard this workbook.
            if execution_failure:
                for index, operation in grouped:
                    if results[index] is None:
                        results[index] = _fail(
                            "operation_execution_failed", execution_failure,
                            index=index, mode=operation["mode"], language=language, path=str(path), persisted=False,
                        )
                continue

            needs_save = any(changed for _, _, changed in staged)
            if needs_save:
                try:
                    _save_workbook_atomic(workbook, path)
                except Exception as exc:
                    for index, result, _ in staged:
                        results[index] = _fail(
                            "workbook_save_failed",
                            _issue(
                                result["sheet"],
                                "persist_success",
                                f"workbook save failed: {exc}",
                                row=result.get("row"),
                                **_file_error_detail(exc, path, "save_workbook"),
                            ),
                            index=index,
                            mode=result["mode"],
                            language=language,
                            path=str(path),
                            sheet=result["sheet"],
                            item=result["item"],
                            persisted=False,
                        )
                    continue
            for index, result, changed in staged:
                record = _row_record(workbook[result["sheet"]], result["row"])
                results[index] = {
                    **result,
                    "index": index,
                    "language": language,
                    "path": str(path),
                    "persisted": changed,
                    "record": record,
                }
                if "frequency" in result:
                    results[index]["frequency"] = record["frequency"]
        finally:
            workbook.close()

    final_results = [
        result
        if result is not None
        else _fail("internal_error", {"reason": "batch result was not produced"}, index=index)
        for index, result in enumerate(results)
    ]
    successes = sum(1 for result in final_results if result["ok"])
    failures = len(final_results) - successes
    if failures:
        reason = "partial_failure" if successes else "batch_failed"
        return _fail(reason, {"successes": successes, "failures": failures}, results=final_results)
    return _ok(results=final_results, successes=successes, failures=0)
